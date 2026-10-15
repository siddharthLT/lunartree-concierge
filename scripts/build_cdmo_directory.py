"""Build concierge/data/cdmo_live_directory.json from the CDMO Live hitlist.

Output holds real attendee PII (names, emails) and is gitignored — it must be
copied to the server out-of-band. Re-run after the hitlist changes; hand-curated
company enrichment (domain, modalities) lives in scripts/company_enrichment.json
and always wins over the heuristics here.
"""
import collections
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'CDMO Live Hitlist (2).xlsx'
ENRICH = Path(__file__).resolve().parent / 'company_enrichment.json'
OUT = ROOT / 'concierge' / 'data' / 'cdmo_live_directory.json'

# Audience: companies that sell into outsourcing. Pharma/biotech outsourcers
# are the buyers (the programs), non-pharma is noise.
KEEP_SEGMENTS = {'CDMO', 'CRDMO', 'CRO', 'Reagents / Formulation Supplier'}
SEGMENT_LABEL = {
    'CDMO': 'CDMO', 'CRDMO': 'CDMO', 'CRO': 'CRO',
    'Reagents / Formulation Supplier': 'Supplier',
}

FREE_MAIL = {'gmail.com', 'yahoo.com', 'outlook.com', 'hotmail.com', 'icloud.com', 'aol.com'}

# keyword -> modality id; first-pass only, enrichment file overrides
MODALITY_KEYWORDS = {
    'api': [r'\bapi\b', r'\bapis\b', r'hpapi', r'active pharmaceutical', r'intermediate', r'drug substance'],
    'adc': [r'\badc', r'conjugat', r'payload', r'linker'],
    'cell': [r'cell therap', r'car-t', r'cell and gene', r'cell & gene', r'\bcgt\b', r'autologous'],
    'gene': [r'gene therap', r'viral vector', r'\baav\b', r'lentivir', r'plasmid', r'cell and gene', r'cell & gene'],
    'editing': [r'gene edit', r'crispr'],
    'mab': [r'monoclonal', r'\bmabs?\b', r'antibod', r'biologic', r'mammalian'],
    'protein': [r'recombinant', r'protein', r'biologic', r'microbial', r'mammalian'],
    'bispecific': [r'bispecific', r'multispecific'],
    'microbiome': [r'microbiome', r'live biotherapeutic'],
    'radio': [r'radiopharm', r'radioligand', r'radiolabel'],
    'small': [r'small molecule', r'\bnce', r'oral solid', r'oral dosage', r'\bosd\b', r'peptide', r'oligo', r'hpapi', r'small-molecule'],
    'device': [r'device', r'combination product', r'autoinjector', r'prefilled syringe', r'packaging'],
    'biosimilar': [r'biosimilar'],
    'digital': [r'digital therap'],
}


def norm(name):
    n = re.sub(r'[^a-z0-9 ]', ' ', (name or '').lower())
    n = re.sub(r'\b(inc|llc|ltd|limited|corp|corporation|co|company|gmbh|ag|sa|plc|group|pvt|private|the|lp|bv|nv|sas|spa|srl|usa|us)\b', ' ', n)
    return re.sub(r'\s+', ' ', n).strip()


def infer_modalities(text):
    t = (text or '').lower()
    return [m for m, pats in MODALITY_KEYWORDS.items() if any(re.search(p, t) for p in pats)]


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True)
    rows = list(wb['Attendees'].iter_rows(values_only=True))
    head = rows[0]
    people = [dict(zip(head, r)) for r in rows[1:] if r[0]]
    ex_rows = list(wb['ExhibitorSponsors'].iter_rows(values_only=True))
    ex_head = ex_rows[0]
    exhibitors = {norm(r[1]): dict(zip(ex_head, r)) for r in ex_rows[1:] if r[1]}
    enrich = json.loads(ENRICH.read_text()) if ENRICH.exists() else {}

    companies = collections.OrderedDict()
    out_people = []
    for p in people:
        if p['Segment'] not in KEEP_SEGMENTS:
            continue
        key = norm(p['company'])
        c = companies.setdefault(key, {
            'name': p['company'], 'segment': SEGMENT_LABEL[p['Segment']],
            'notes': p.get('Sub-type / Notes') or '', 'domains': collections.Counter(),
        })
        email = (p.get('Work Email (Apollo)') or '').strip().lower()
        if email and p.get('Email Status') != 'unavailable' and '@' in email:
            d = email.split('@')[1]
            if d not in FREE_MAIL:
                c['domains'][d] += 1
        else:
            email = ''
        out_people.append({
            'id': str(int(p['profile_id'])) if p.get('profile_id') else f"x{len(out_people)}",
            'name': p['name'].strip(), 'title': (p.get('job_title') or '').strip(),
            'company': key, 'email': email,
        })

    out_companies = {}
    for key, c in companies.items():
        ex = exhibitors.get(key) or next((v for k, v in exhibitors.items() if k and (k in key or key in k)), None)
        blob = ' '.join(filter(None, [c['notes'], (ex or {}).get('Sub-type / Notes'), (ex or {}).get('summary')]))
        e = enrich.get(key, {})
        out_companies[key] = {
            'name': e.get('name') or c['name'],
            'segment': c['segment'],
            'domain': e.get('domain') or (c['domains'].most_common(1)[0][0] if c['domains'] else ''),
            'modalities': e.get('modalities') if 'modalities' in e else infer_modalities(blob),
            'exhibitor': bool(ex),
        }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({'companies': out_companies, 'people': out_people}, separators=(',', ':')))
    n_dom = sum(1 for c in out_companies.values() if c['domain'])
    n_mod = sum(1 for c in out_companies.values() if c['modalities'])
    print(f"{len(out_people)} people, {len(out_companies)} companies, {n_dom} with domain, {n_mod} with modalities", file=sys.stderr)


if __name__ == '__main__':
    main()
